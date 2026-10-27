import ast
import math
import os

import pytest

import design as D
from model import checks, geometry as G, hydro, variants, weights

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


@pytest.fixture(scope='module')
def r():
    return hydro.design()


def test_hydro_displacement_equals_weight(r):
    assert abs(r['disp'] - r['W']) < 1.0


def test_hydro_lcb_equals_lcg(r):
    assert abs(r['lcb'] - r['G']) < 0.1


def test_deck_parallel_to_waterline_at_design_position(r):
    for X in (0, D.LOA / 2, D.LOA):
        assert abs(r['deck'](X) - (r['d0'] - r['k'] * X) - D.FREEBOARD) < 1e-9


def test_other_battery_positions_balance():
    for b in D.BAT_SLIDE:
        rr = hydro.at(b)
        assert abs(rr['disp'] - rr['W']) < 1.0
        assert abs(rr['lcb'] - rr['G']) < 0.1


def test_sponson_width_at_transom_is_design_aft_width():
    assert abs(G.sponson_width(0) - D.AFT_W) < 1e-9
    assert abs(G.sponson_width(D.AFT_TAPER_X) - D.SPONSON_W) < 1e-9
    assert abs(G.sponson_width(D.LOA) - D.TIP_W) < 1e-9


def test_taper_symmetric_about_sponson_centerline():
    for X in (0, 100, 200, 300):
        mid = (G.inner_wall_y(X) + G.outer_wall_y(X)) / 2
        assert abs(mid - G.YS) < 1e-9


def test_derived_values():
    assert abs(G.BOW0 - 706.1) < 0.05
    assert abs(G.shaft_z(D.CV_X) - D.PROP_Z) < 1e-9
    assert abs(G.shaft_z(G.KEEL_X)) < 1e-9
    assert abs(G.SKEG_X0 - (G.KEEL_X + D.SKEG_FWD_OF_KEEL)) < 1e-9
    # aft face axis height = SKIN_T + body radius / cos(axis) + clearance
    assert abs(G.F[1] - (D.SKIN_T + G.BODY_D / 2 / math.cos(math.radians(D.MOTOR_AXIS_DEG)) + D.MOTOR_FLOOR_CLR)) < 1e-9
    assert G.MOT_LOW > D.SKIN_T


def test_blade_retracted_is_rotated_up_by_the_sector():
    ext, ret = G.blade_outline(True), G.blade_outline(False)
    assert min(p[1] for p in ext) < D.BLADE_VARIANTS[D.BLADE_VARIANT]['hinge_z']
    assert max(p[1] for p in ret) > D.BLADE_VARIANTS[D.BLADE_VARIANT]['hinge_z']


def test_weights_have_unweighed_items_and_total(r):
    assert all(m is None for _, _, m, _ in weights.UNWEIGHED)
    assert abs(sum(i[2] for i in r['items']) - r['W']) < 1e-9


def test_checks_run_and_report(r):
    res = checks.run_all(r)
    assert len(res) >= 10
    for name, ok, msg in res:
        assert isinstance(name, str) and isinstance(ok, (bool,)) and msg
    text = checks.report(res)
    assert 'CHECKS REPORT' in text


def test_check_functions_individually(r):
    assert checks.c1(r)[1] is True
    assert checks.c7(r)[1] is True
    assert checks.c8(r)[1] == (r['bow'] >= D.BOW_DRAFT_MIN)


def test_no_numeric_literals_outside_design():
    allowed = {0, 1, 2, 1000}
    files = ['build.py'] + [os.path.join('model', f) for f in ('geometry.py', 'weights.py', 'hydro.py', 'checks.py', 'variants.py', 'schedules.py', 'study.py')]
    bad = []
    for f in files:
        tree = ast.parse(open(os.path.join(ROOT, f)).read())
        docs = {id(n.body[0].value) for n in ast.walk(tree) if isinstance(n, (ast.Module, ast.FunctionDef, ast.ClassDef)) and n.body
                and isinstance(n.body[0], ast.Expr) and isinstance(getattr(n.body[0], 'value', None), ast.Constant)}
        for n in ast.walk(tree):
            if isinstance(n, ast.Constant) and isinstance(n.value, (int, float)) and not isinstance(n.value, bool):
                if n.value not in allowed:
                    bad.append((f, n.lineno, n.value))
    assert not bad, bad


# ---- reverse daggerboard variant and the override mechanism -------------------------------------

def test_reverse_tip_positions():
    b = D.BLADE_VARIANTS['reverse']
    a = math.radians(b['sector_deg'])
    ext = G.blade_outline(True, 'reverse')
    # upper edge tip: forward along the keel line, X hinge + R, at the hinge Z
    assert abs(ext[1][0] - (b['hinge_x'] + b['R'])) < 1e-9 and abs(ext[1][1] - b['hinge_z']) < 1e-9
    assert abs(ext[1][0] - 800.0) < 1e-9 and abs(ext[1][1] - 6.0) < 1e-9
    # lower edge tip: forward and down by the sector angle
    assert abs(ext[-1][0] - (b['hinge_x'] + b['R'] * math.cos(a))) < 1e-9
    assert abs(ext[-1][1] - (b['hinge_z'] - b['R'] * math.sin(a))) < 1e-9
    assert abs(G.blade_depth('reverse') - (b['R'] * math.sin(a) - b['hinge_z'])) < 1e-9
    ret = G.blade_outline(False, 'reverse')
    # retracted: rotated up by the sector angle, lying between forward along the keel and forward 15 deg up
    assert abs(ret[-1][0] - 800.0) < 1e-9 and abs(ret[-1][1] - b['hinge_z']) < 1e-9
    assert abs(ret[1][0] - (b['hinge_x'] + b['R'] * math.cos(a))) < 1e-9
    assert abs(ret[1][1] - (b['hinge_z'] + b['R'] * math.sin(a))) < 1e-9


def test_blade_area_and_centroid_formula():
    for key in ('trail', 'hang', 'reverse'):
        b = D.BLADE_VARIANTS[key]
        a = math.radians(b['sector_deg'])
        assert abs(G.blade_area(key) - b['R'] ** 2 * a / 2) < 1e-6
        half = a / 2
        d = 2 * b['R'] * math.sin(half) / (3 * half)
        cx, cz = G.blade_centroid(True, key)
        assert abs(math.hypot(cx - b['hinge_x'], cz - b['hinge_z']) - d) < 1e-9
        sgn = 1 if b['dir'] == 'fwd' else -1
        assert abs((cx - b['hinge_x']) - sgn * d * math.cos(half)) < 1e-9
        assert abs(G.blade_center_from_bow(key) - (D.LOA - cx)) < 1e-9


def test_reverse_area_numbers():
    assert abs(G.blade_area('reverse') - 380.0 ** 2 * math.radians(15.0) / 2) < 1e-6
    assert abs(G.blade_depth('reverse') - (380.0 * math.sin(math.radians(15.0)) - 6.0)) < 1e-9


def test_override_context_restores_and_runs_a_variant():
    sel, keys = D.BLADE_VARIANT, set(D.BLADE_VARIANTS)
    with variants.use('reverse', dict(sector_deg=20.0), name='reverse 20 test') as p:
        assert D.BLADE_VARIANT == 'reverse 20 test' and p['sector_deg'] == 20.0
        assert abs(G.blade_area() - 380.0 ** 2 * math.radians(20.0) / 2) < 1e-6
        rr = hydro.design()
        assert abs(rr['disp'] - rr['W']) < 1.0
    assert D.BLADE_VARIANT == sel and set(D.BLADE_VARIANTS) == keys
    assert D.BLADE_VARIANTS['reverse']['sector_deg'] == 15.0


def test_reverse_housing_is_a_bow_cap_over_the_stowed_blade():
    with variants.use('reverse'):
        rr = hydro.design()
        zd = rr['deck']
        h = G.housing(zd)
        assert h['kind'] == 'cap'
        pts = G.stowed_above(zd)
        assert abs(h['x0'] - min(p[0] for p in pts)) < 1e-9 and abs(h['x1'] - max(p[0] for p in pts)) < 1e-9
        assert abs(h['ztop'] - (max(max(p[1] for p in G.blade_outline(False)), G.blade_servo_box(zd).z1) + D.CAP_CLR)) < 1e-9
        # at the cap edges the stowed blade crosses the deck underside
        for x in (h['x0'], h['x1']):
            assert min(abs(p[1] - (zd(p[0]) - D.SKIN_T)) for p in pts if abs(p[0] - x) < 1e-9) < 1e-6
        names = [n for _, n, _, _ in rr['items']]
        assert any('Bow cap' in n for n in names) and any('in the bow cap' in n for n in names)
        blades = sum(m for g, _, m, _ in rr['items'] if g == 'Blades')
        per = G.blade_area() * D.SKIN_T * D.BALSA_RHO / 1000 + 2 * G.blade_area() / 1e6 * D.GLASS_OUT
        assert abs(blades - (2 * per + D.TRUNK_MASS + D.TOWER_MASS + D.BLADE_MECH_MASS)) < 1e-9
        res = {n.split()[0]: ok for n, ok, _ in checks.run_all(rr)}
        assert res['9'] and res['9b']


def test_sweep_runs_every_variant_at_both_battery_positions():
    rows, runs = variants.sweep()
    assert len(rows) == len(D.BLADE_SWEEP) * len(D.SWEEP_BAT_X)
    assert {r['battery_x'] for r in rows} == set(D.SWEEP_BAT_X)
    assert D.BLADE_VARIANT == 'reverse25'
    row = [r for r in rows if r['variant'] == 'reverse 25 deg R375' and r['battery_x'] == D.DESIGN_BAT_X][0]
    assert abs(row['total_g'] - hydro.design()['W']) < 1e-9
    hang = [r for r in rows if r['variant'] == 'hang R120' and r['battery_x'] == D.DESIGN_BAT_X][0]
    with variants.use('hang'):
        assert abs(hang['total_g'] - hydro.design()['W']) < 1e-9


# ---- reverse 25 deg daggerboard as the default, servo on its side, receiver and battery offsets --------------

def test_default_blade_is_reverse25():
    b = D.BLADE_VARIANTS[D.BLADE_VARIANT]
    assert D.BLADE_VARIANT == 'reverse25'
    assert (b['hinge_x'], b['hinge_z'], b['R'], b['sector_deg']) == (420.0, 6.0, 375.0, 25.0)
    assert b['dir'] == 'fwd' and b['side'] == 'tunnel' and b['housing'] == 'cap'
    assert (b['trunk_x0'], b['trunk_x1'], b['trunk_w']) == (420.0, 800.0, 12.0)
    ext, ret = G.blade_outline(True), G.blade_outline(False)
    assert abs(ext[1][0] - (800.0 - D.TRUNK_TIP_CLR)) < 1e-9 and abs(ext[1][1] - 6.0) < 1e-9
    assert abs(G.blade_depth() - (375.0 * math.sin(math.radians(25.0)) - 6.0)) < 1e-9
    assert abs(max(p[1] for p in ret) - (6.0 + 375.0 * math.sin(math.radians(25.0)))) < 1e-9


def test_reverse25_blade_plane_and_trunk_on_the_tunnel_side():
    ty0, ty1 = G.trunk_y()
    assert abs(ty0 - (G.HT + D.SKIN_T)) < 1e-9
    assert abs((ty1 - ty0) - (12.0 + D.SKIN_T)) < 1e-9
    assert abs(G.blade_y() - (G.HT + D.SKIN_T + 6.0)) < 1e-9
    assert ty0 < G.blade_y() < ty1


def test_reverse25_cap_over_the_stowed_tip(r):
    h = G.housing(r['deck'])
    tip = 6.0 + 375.0 * math.sin(math.radians(25.0))
    assert h['kind'] == 'cap' and abs(h['ztop'] - (tip + D.CAP_CLR)) < 1e-9
    assert h['h0'] > 0 and h['h1'] > 0
    assert 0 < h['x0'] < h['x1'] <= D.LOA
    names = [n for _, n, _, _ in r['items']]
    assert any(n.startswith('Bow cap') for n in names) and not any('Tower' in n for n in names)


def test_steering_servo_lies_on_its_side_on_the_roof():
    b = G.steer_servo_box()
    assert (D.STEER_SERVO['h'], D.STEER_SERVO['l'], D.STEER_SERVO['w']) == (20.0, 40.0, 40.5)
    assert abs((b.z1 - b.z0) - 20.0) < 1e-9 and abs(b.z0 - G.ROOF_Z1) < 1e-9
    assert abs((b.x1 - b.x0) - 40.0) < 1e-9 and abs((b.y1 - b.y0) - 40.5) < 1e-9
    assert (b.x0, b.x1) == (20.0, 60.0) and abs(b.y0 + b.y1) < 1e-9
    assert b.y1 <= G.HT


def test_receiver_moved_forward_and_clears_the_escs():
    rx = G.rx_box()
    assert (rx.x0, rx.x1) == (345.0, 392.0)
    for e in G.esc_boxes():
        assert not (min(rx.x1, e.x1) - max(rx.x0, e.x0) > 0 and min(rx.y1, e.y1) - max(rx.y0, e.y0) > 0)


def test_battery_offset_outboard_clears_the_trunk():
    assert D.BAT_Y_OFFSET == 5.0
    assert abs(G.BAT_YC - 97.075) < 1e-9 and abs(G.YS - 92.075) < 1e-9
    port, stbd = G.battery_boxes(D.DESIGN_BAT_X)
    assert abs(port.y0 + port.y1 - 2 * 97.075) < 1e-9 and abs(stbd.y0 + stbd.y1 + 2 * 97.075) < 1e-9
    _, ty1 = G.trunk_y()
    assert port.y0 > ty1                                   # clear of the 12 mm trunk
    assert G.YS - D.BAT['w'] / 2 < ty1                     # and it would not be without the offset
    assert port.y1 <= float(G.outer_wall_y(G.battery_xrange(D.BAT_SLIDE[1])[1])) - D.SKIN_T


def test_known_check_state(r):
    """Current results: every check passes except 21, the three owner-review conflicts reported on A-502 (F5 tunnel piece, F7 rib)."""
    res = {n.split()[0]: ok for n, ok, _ in checks.run_all(r)}
    assert {k for k, ok in res.items() if not ok} == {'21'}
    assert len(res) == 22 and res['19'] and res['20']
    assert res['6'] and res['15'] and res['17'] and res['18'] and res['13'] and res['9'] and res['9b'] and res['3'] and res['11']


# ---- v9.1: design position 475, hatches HB and HC, linkage, shaft, density, checks 15 to 18 ---------------------

def test_revision_and_design_position():
    assert D.REV == 'v9.2'
    assert D.DESIGN_BAT_X == 475.0 and D.BAT_SLIDE == (475.0, 590.0) and D.ALT_BAT_X == (520.0, 590.0)
    assert D.SWEEP_BAT_X == (475.0, 520.0)
    assert hydro.design()['bat_x'] == 475.0


def test_deck_is_level_at_475_and_alternates_are_slightly_off(r):
    assert abs(r['deck_off_level']) < 1e-9
    assert abs(hydro.at(520.0, r)['deck_off_level']) > 0.1


def test_hatch_hb_is_centered_on_the_shifted_battery_and_hc_covers_the_escs():
    hb = D.HATCHES['HB']
    assert (hb['x0'], hb['x1'], hb['w']) == (425.0, 582.0, 48.0)
    op = G.hatch_open('HB', G.BAT_YC)
    assert abs((op['y0'] + op['y1']) / 2 - G.BAT_YC) < 1e-9 and abs((op['y1'] - op['y0']) - 48.0) < 1e-9
    assert abs(G.hatch_open('HB', -G.BAT_YC)['yc'] + G.BAT_YC) < 1e-9
    hc = D.HATCHES['HC']
    assert (hc['x0'], hc['x1'], hc['w']) == (265.0, 400.0, 94.0)
    for e in G.esc_boxes():
        assert hc['x0'] <= e.x0 and e.x1 <= hc['x1']


def test_receiver_is_forward_of_the_escs_inside_hc():
    rx = G.rx_box()
    assert (rx.x0, rx.x1) == (345.0, 392.0)
    hc = G.hatch_open('HC', 0.0)
    assert hc['x0'] <= rx.x0 and rx.x1 <= hc['x1']
    for e in G.esc_boxes():
        assert e.x1 <= rx.x0                                   # forward of the ESCs, no overlap in X


def test_steering_linkage_geometry():
    assert abs(G.ZR - (G.ROOF_Z1 + 10.0)) < 1e-9
    assert G.ROD_X1 == G.POSTX and abs(G.ROD_L - (G.ROD_X0 - G.POSTX)) < 1e-9 and G.ROD_L > 0
    assert abs(G.TIE_L - 2 * G.YS) < 1e-9                       # post spacing, horns parallel
    assert abs(G.TIE_Y1 - (G.YS + D.NOZ_HORN_ARM)) < 1e-9
    assert (D.TIE_BAR, D.NOZ_HORN_ARM, D.PINTLE_D) == (4.0, 20.0, 3.0)
    # the transom hole sits in the center transom closure, above the roof and below the aft deck
    r0 = hydro.design()
    ty = G.HT + float(G.inner_off(0))
    assert G.ROD_Y + D.BELLOWS_OD / 2 < ty
    assert G.ROOF_Z1 < G.ZR - D.ROD_HOLE_D / 2 and G.ZR + D.ROD_HOLE_D / 2 < r0['deck'](0) - D.SKIN_T
    # pintle: pin from the nozzle bottom into the shoe bushing, shorter than the shoe thickness
    assert abs(G.PINTLE_L - (D.SHOE_CLR + D.PINTLE_ENGAGE)) < 1e-9 and D.PINTLE_ENGAGE <= D.SHOE_T and D.BUSHING_OD > D.PINTLE_D


def test_post_required_length_and_check_18(r, monkeypatch):
    need = G.post_required_length(r['deck'])
    assert abs(need - (r['deck'](G.POSTX) - D.AFT_DECK_T - G.NOZ_ZTOP)) < 1e-9 and 94.0 < need < 96.0
    assert not hasattr(D, 'NOZ_POST_TUBE_L') and not hasattr(D, 'NOZ_POST_L') and D.NOZ_POST_STOCK_L == 300.0
    name, ok, msg = checks.c18(r)
    assert ok and '%.1f' % need in msg and 'cut' in msg
    monkeypatch.setattr(D, 'NOZ_POST_STOCK_L', 50.0)
    assert not checks.c18(r)[1]                               # the check compares to the cut stock, here a 50 mm piece


def test_parts_has_post_rod_and_tube_with_no_price_or_mass():
    from model import parts as P
    for k in ('post_rod', 'post_tube'):
        p = P.get(k)
        assert p.price is None and p.mass is None and p.qty == 1 and '300 mm' in p.spec and 'Source' in p.note


def test_shaft_length_is_derived():
    assert not hasattr(D, 'SHAFT_LEN') and (D.SHAFT_STUB, D.SHAFT_MARGIN) == (10.0, 5.0)
    prop = (D.CV_X - G.POSTX) + D.SHAFT_MARGIN
    motor = (G.MOT_X0 - D.CV_X) / math.cos(math.radians(D.MOTOR_AXIS_DEG)) + D.SHAFT_STUB
    assert abs(G.SHAFT_LEN - (prop + motor)) < 1e-9 and abs(G.SHAFT_LEN - 373.3) < 0.1
    assert G.SHAFT_LEN < 390.0
    assert G.POSTX < G.SHAFT_CG_X < G.MOT_X0
    items = {n: m for _, n, m, _ in hydro.design()['items']}
    shaft = [m for n, m in items.items() if n.startswith('Shafts')][0]
    assert abs(shaft - 2 * math.pi * (D.SHAFT_D / 2) ** 2 * G.SHAFT_LEN / 1000 * D.STEEL_RHO) < 1e-9


def test_density_sensitivity(r):
    assert D.BALSA_RHO_HI == 0.160
    (rho0, r0), (rho1, r1) = hydro.density_cases(r)
    assert rho0 == D.BALSA_RHO and r0 is r and rho1 == 0.160
    assert r1['W'] > r['W'] and abs(r1['disp'] - r1['W']) < 1.0 and r1['rho'] == 0.160
    # only the balsa mass changes: the rest of the item list is identical
    d = sum(m for _, _, m, _ in r1['items']) - sum(m for _, _, m, _ in r['items'])
    balsa = sum(m for g, n, m, _ in r['items'] if g in ('Hull skin',)) + sum(m for g, n, m, _ in r['items'] if n.startswith('Daggerboards'))
    assert d > 0 and d < balsa * (0.160 / D.BALSA_RHO - 1) * 1.1
    assert abs(r1['over'] - (r1['d0'] - G.NOZ_ZTOP)) < 1e-9
    assert abs(hydro.at_density(D.BALSA_RHO, r)['W'] - r['W']) < 1e-9


def test_check_6_evaluates_at_the_design_position(r):
    name, ok, msg = checks.c6(r)
    assert ok == (r['over'] - D.TROUGH >= D.TROUGH_MARGIN_MIN) and 'batteries at 475' in msg and checks.level(name, ok) == 'PASS'
    with variants.use('reverse25'):
        alt = hydro.at(590.0, hydro.design())
    nm, ok2, _ = checks.c6(alt)
    assert not ok2 and checks.level(nm, ok2) == 'WARN'            # still a warning at 590, not a failure


def test_check_15_only_tests_the_open_part_of_the_trunk(r, monkeypatch):
    zd = r['deck']
    above = G.stowed_above(zd)
    open_x0 = min(p[0] for p in above)
    assert checks.c15(r)[1] and '(5.1 clear in X)' in checks.c15(r)[2]          # HB x1 582 plus the 6 flange is 588, the open trunk starts at 593.1
    monkeypatch.setitem(D.HATCHES, 'HB', dict(D.HATCHES['HB'], x1=open_x0 - D.FLANGE + 2.5))        # the flange reaches 2.5 into the open part
    assert checks.c15(r)[1] is False and 'open trunk by 2.5 in X' in checks.c15(r)[2]
    monkeypatch.setitem(D.HATCHES, 'HB', dict(D.HATCHES['HB'], x1=open_x0 - D.FLANGE - 0.5))
    assert checks.c15(r)[1]
    # a hatch over the closed part of the trunk (Y overlaps the trunk, X before the open part) passes
    monkeypatch.setitem(D.HATCHES, 'HB', dict(x0=430.0, x1=560.0, w=60.0, zone='sponson'))
    ty0, ty1 = G.trunk_y()
    assert any(h['y0'] < ty1 and h['y1'] > ty0 for h in G.hatch_rects(flange=True) if h['name'] == 'HB' and h['yc'] > 0)
    assert checks.c15(r)[1]


def test_check_16_battery_at_the_service_position(r):
    assert D.BAT_SERVICE_X == 504.0 and D.BAT_SLIDE[0] <= D.BAT_SERVICE_X <= D.BAT_SLIDE[1]
    name, ok, msg = checks.c16(r)
    assert ok                                                 # the 155 battery at 504 spans X 426.5 to 581.5, HB opens 425 to 582
    items = G.lift_items([D.BAT_SERVICE_X])
    assert [fp['name'] for fp in items if 'bat_x' in fp] == ['battery at 504']                # not tested at every slide position
    fits = {fp['name']: checks._fit(fp, G.hatch_open(fp['hatch'], (fp['y0'] + fp['y1']) / 2)) for fp in items}
    assert fits['motor and jacket'].inside and fits['ESC port'].inside and fits['ESC stbd'].inside and fits['steering servo'].inside
    assert fits['receiver'].inside                            # HC now covers the receiver
    b = fits['battery at 504']
    assert b.inside and b.outx == 0 and b.outy == 0
    assert 'battery at 475' not in msg and 'battery at 590' not in msg


def test_hb_flange_clearance_to_the_open_trunk():
    hb = [h for h in G.hatch_rects(flange=True) if h['name'] == 'HB' and h['yc'] > 0][0]
    assert abs(hb['x1'] - 588.0) < 1e-9


def test_check_17_bumper_margins(r, monkeypatch):
    name, ok, msg = checks.c17(r)
    sweep = G.nozzle_sweep()
    end = sweep[-1]
    assert abs(end['deg'] - 90.0) < 1e-9
    assert abs(end['aft'] - (-G.POSTX + D.NOZ_OD / 2)) < 1e-9 and abs(end['out'] - (G.YS + D.NOZ_L / 2)) < 1e-9     # a square plan footprint at 90 deg
    corner = math.hypot(D.NOZ_L / 2, D.NOZ_OD / 2)
    mid = [x for x in sweep if abs(x['deg'] - 45.0) < 1e-9][0]
    assert abs(mid['out'] - (G.YS + corner)) < 1e-9 and abs(mid['aft'] - (-G.POSTX + corner)) < 1e-9
    assert D.AFT_DECK_W == 285.0
    assert ok and abs((D.AFT_DECK_W / 2 - (G.YS + corner)) - 8.0) < 0.01 and 'outboard margin 8.0 at 45.0 deg' in msg
    assert 'aft margin 19.2 at 45.0 deg' in msg and 'margin 31.6' in msg and 'margin 20.4' in msg
    monkeypatch.setattr(D, 'AFT_DECK_W', 270.0)
    assert not checks.c17(r)[1]                              # the old width fails the 5 mm margin


def test_every_sheet_carries_the_revision(r):
    from sheets import a101, a104, a301, a601, a004
    rows, _ = variants.sweep()
    alts = [hydro.at(x, r) for x in D.ALT_BAT_X]
    results = checks.run_all(r)
    for sh in (a101.build(r, alts), a104.build(r), a301.build(r, alts), a601.build(r, alts, results), a004.build(rows)):
        assert sh.rev == D.REV == 'v9.2'
        assert 'v9.2, issued for review' in sh.svg()
    svg = a601.build(r, alts, results).svg()
    assert '0.160 g/cm3 (10 lb/ft3)' in svg and 'DENSITY SENSITIVITY' in svg
    assert 'Pushrod, 2-56 rod' in a101.build(r, alts).svg()


# ---- schedule sheets, open items, parts -----------------------------------------------------------

def test_every_sheet_builds(r):
    from sheets import a000, a101, a103, a104, a301, a401, a501, a502, a503, a601, a004
    alts = [hydro.at(x, r) for x in D.ALT_BAT_X]
    sweep, _ = variants.sweep()
    shs = a000.build(r, [a000.SELF + (a000.CONTENTS,)]) + [a101.build(r, alts), a103.build(r), a104.build(r), a301.build(r, alts),
           a401.build(r, alts), a501.build(r), a502.build(r), a503.build(r), a601.build(r, alts, checks.run_all(r)), a004.build(sweep)]
    assert len(shs) == 12 and all(s.svg().startswith('<svg') for s in shs)


def test_open_items_find_every_tag():
    from model import schedules as S
    tags = {i['tag'] for i in S.open_items()}
    assert tags == set(S.TAGS)


def test_parts_masses_match_design_and_weight_total(r):
    from model import parts as P
    assert P.mass('motor') == D.MOTOR_MASS and P.mass('esc') == D.ESC['mass'] and P.mass('battery') == D.BAT['mass']
    assert P.qty('jacket') == (2 if D.JACKETS else 0)
    assert abs(sum(i[2] for i in r['items']) - r['W']) < 1e-9


def test_build_list_writes_workbook(tmp_path, r):
    import openpyxl
    from tools import build_list
    path = build_list.main(str(tmp_path), r)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[:2] == ['Build List', 'Channels and Wiring']
    for ws in wb:
        for row in ws.iter_rows(values_only=True):
            assert not any(isinstance(v, str) and (chr(0x2014) in v or chr(0x2013) in v) for v in row)


# ---- v9.2 owner review ---------------------------------------------------------------------------------------

def test_v92_constants():
    assert D.PROP_D == 45.0 and D.PROP_D_ALT == 50.0 and D.NOZ_OD == 60.0 and D.NOZ_L == 60.0 and D.NOZ_OD_ALT == 65.0
    assert D.BILGE_R == 15.0 and D.SKEG_ROOT_R == 8.0 and D.SKEG_EDGE_R == 6.0
    assert D.TRUNK_TIP_CLR == 5.0 and D.LINE_R == 200.0 and D.BLADE_SERVO_BAY == (600.0, 650.0)
    assert (D.BLADE_SERVO['l'], D.BLADE_SERVO['w'], D.BLADE_SERVO['h']) == (40.0, 40.5, 20.0)
    assert not hasattr(D, 'BLADE_BOX')
    assert D.MOTOR_KV == 1200.0 and D.FLOOR_FRACTION == 0.08 and dict(D.PACK_V) == {'3S': 11.1, '2S': 7.4}
    assert D.WIRE_HOLE_D == 10.0 and D.FRAME_KIND == {1: 'solid', 2: 'solid', 3: 'solid_shaft', 9: 'solid_notch'}
    assert D.TUNNEL_KIND == {5: 'solid_grommet'}


def test_tip_clearance_and_servo_in_the_cap(r):
    assert abs(G.tip_to_trunk_end() - D.TRUNK_TIP_CLR) < 1e-9
    zd = r['deck']
    sb = G.blade_servo_box(zd)
    assert abs((sb.z1 - sb.z0) - 20.0) < 1e-9 and abs((sb.x1 - sb.x0) - 40.0) < 1e-9 and abs((sb.y1 - sb.y0) - 40.5) < 1e-9
    assert D.BLADE_SERVO_BAY[0] <= sb.x0 and sb.x1 <= D.BLADE_SERVO_BAY[1]
    assert sb.y0 >= G.trunk_y()[1]
    h = G.housing(zd)
    assert abs(h['ztop'] - (max(h['z_blade'], sb.z1) + D.CAP_CLR)) < 1e-9
    res = {n.split()[0]: ok for n, ok, _ in checks.run_all(r)}
    assert res['19'] and res['20'] and not res['21']


def test_rounds_reduce_waterplane_and_raise_draft(r):
    assert abs(G.round_area(D.BILGE_R, D.BILGE_R) - D.BILGE_R ** 2 * (1 - math.pi / 4)) < 1e-9
    assert G.round_area(D.BILGE_R, 0.0) == 0.0
    sq = hydro.design(rounds=False)
    assert r['d0'] > sq['d0'] and abs(r['W'] - sq['W']) < 10.0     # no weight change by the rounds; only the deck skin follows the deeper draft
    e = hydro.round_effect(r)
    assert e['d0'] > 0 and e['bow'] > 0
    assert G.SKEG_VOL > G.SKEG_VOL_BASE and G.SKEG_VOL_FILLET > 0


def test_frames_and_compartments(r):
    from model import schedules as S
    kinds = [f['kind'] for f in S.frames(r['deck'])]
    assert kinds[0] == 'solid' and kinds[1] == 'solid' and kinds[2] == 'solid_shaft' and kinds[8] == 'solid_notch'
    assert all(k == 'ring' for i, k in enumerate(kinds) if i + 1 in (4, 6, 7))


def test_study_alternate_restores_state(r):
    from model import study
    before = (D.PROP_D, D.NOZ_OD, hydro.design()['d0'])
    rows = study.rows(r, checks.run_all(r))
    assert (D.PROP_D, D.NOZ_OD, hydro.design()['d0']) == before
    assert len(rows) == 2


def test_slow_speed_numbers():
    from model import schedules as S
    t = repr(S.slow_speed())
    assert '13320' in t.replace(',', '') and '8880' in t.replace(',', '')


def test_a004_is_not_in_the_set():
    src = open(os.path.join(ROOT, 'build.py')).read()
    assert 'a004' not in src
    assert os.path.exists(os.path.join(ROOT, 'sheets', 'a004.py'))
